"""
Générateur de questionnaire .docx à partir d'une note de cadrage et d'un cadre de résultats.
L'IA analyse les documents d'entrée et produit un questionnaire structuré
au format standard SISTA (tables à 5 colonnes).
"""
import json
import io
import re
from docx import Document
from docx.shared import Pt, Cm, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
import openpyxl

from . import ai_agent


def extract_note_cadrage(file_bytes: bytes) -> str:
    """Extraire le contenu textuel d'une note de cadrage (.docx)."""
    doc = Document(io.BytesIO(file_bytes))
    lines = []
    for p in doc.paragraphs:
        txt = p.text.strip()
        if txt:
            style = p.style.name if p.style else ""
            if "Heading" in style:
                lines.append(f"\n## {txt}")
            else:
                lines.append(txt)
    for i, table in enumerate(doc.tables):
        lines.append(f"\n[Tableau {i+1}]")
        for row in table.rows:
            cells = [c.text.strip()[:100] for c in row.cells]
            lines.append(" | ".join(cells))
    return "\n".join(lines)


def extract_cadre_resultats(file_bytes: bytes, filename: str) -> str:
    """Extraire le contenu d'un cadre de résultats (.xlsx ou .docx)."""
    if filename.lower().endswith((".xlsx", ".xls")):
        wb = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True)
        lines = []
        for ws in wb.worksheets:
            lines.append(f"\n=== Feuille: {ws.title} ===")
            for row in ws.iter_rows(values_only=True):
                vals = [str(v)[:80] if v else "" for v in row]
                if any(vals):
                    lines.append(" | ".join(vals))
        return "\n".join(lines)
    elif filename.lower().endswith((".docx", ".doc")):
        return extract_note_cadrage(file_bytes)
    else:
        return file_bytes.decode("utf-8", errors="replace")


def generate_questionnaire_data(
    note_cadrage_text: str,
    cadre_resultats_text: str,
    api_key: str,
    api: str = "api1",
) -> dict:
    """Appeler l'IA pour générer la structure du questionnaire en JSON."""

    system_prompt = """Tu es un expert en conception de questionnaires d'enquêtes quantitatives pour des projets de développement (HCR, UNICEF, Banque Mondiale, etc.).

À partir de la note de cadrage (méthodologie) et du cadre de résultats (matrice des indicateurs) fournis, tu dois concevoir un questionnaire quantitatif complet et professionnel.

Le questionnaire doit suivre cette structure standard :

1. En-tête : titre du projet, type de questionnaire, population cible, zones, version
2. Instructions générales pour les enquêteurs
3. Parties numérotées, chacune contenant des questions sous forme de tableau à 5 colonnes :
   - N° (numérotation hiérarchique : 1.1, 1.2, etc.)
   - Question (texte complet de la question)
   - Modalités de réponse (choix numérotés, ou format de réponse)
   - Filtre / Instruction (consignes pour l'enquêteur, sauts conditionnels)
   - Indicateur mesuré (code de l'indicateur du cadre de résultats, ou "—")

Les parties obligatoires sont :
- Partie 0 : Identification (ID questionnaire, enquêteur, date, lieu)
- Partie 1 : Consentement libre et éclairé
- Partie 2 : Caractéristiques sociodémographiques
- Parties 3+ : Sections thématiques basées sur les indicateurs du cadre de résultats
- Dernière partie : Clôture de l'entretien

Règles :
- Chaque indicateur du cadre de résultats doit être couvert par au moins une question
- Les modalités de réponse doivent être précises et numérotées
- Ajouter des filtres/sauts logiques quand nécessaire
- Respecter la sensibilité des sujets (VBG, protection, etc.)
- Le questionnaire doit être en français

Réponds UNIQUEMENT avec un JSON valide (pas de markdown, pas de commentaires) avec cette structure :
{
  "header": {
    "organization": "Nom de l'organisation",
    "project_title": "Titre du projet",
    "questionnaire_type": "Type de questionnaire",
    "target_population": "Population cible",
    "zones": "Zones géographiques",
    "version": "Version 1"
  },
  "instructions": ["instruction 1", "instruction 2", ...],
  "parts": [
    {
      "number": 0,
      "title": "Identification de l'enquête",
      "description": "Description optionnelle ou texte introductif",
      "sub_sections": [],
      "questions": [
        {
          "number": "0.1",
          "question": "Texte de la question",
          "modalities": "1. Option A  2. Option B  3. Option C",
          "filter": "Instruction pour l'enquêteur ou —",
          "indicator": "Code indicateur ou —"
        }
      ]
    }
  ],
  "closure_script": "Texte de clôture de l'entretien",
  "indicator_mapping": [
    {"indicator_code": "Ind.1a", "questions": ["3.1", "3.2"]}
  ]
}"""

    user_prompt = f"""Voici les documents d'entrée :

=== NOTE DE CADRAGE (MÉTHODOLOGIE) ===
{note_cadrage_text[:12000]}

=== CADRE DE RÉSULTATS (MATRICE DES INDICATEURS) ===
{cadre_resultats_text[:8000]}

Génère le questionnaire complet en JSON selon le format demandé."""

    api_config = ai_agent.API_CONFIG.get(api, ai_agent.API_CONFIG["api1"])
    model = api_config["model_smart"]

    response_text = ai_agent._call_with_retry(
        api_key=api_key,
        api=api,
        model=model,
        system_prompt=system_prompt,
        user_prompt=user_prompt,
        max_tokens=8000,
    )

    # Extraire le JSON
    data = ai_agent._extract_json(response_text)
    if not data:
        raise ValueError("L'IA n'a pas retourné un JSON valide pour le questionnaire.")
    return data


def build_questionnaire_docx(data: dict) -> bytes:
    """Construire le fichier .docx du questionnaire à partir du JSON structuré."""
    doc = Document()

    # -- Styles de base --
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)
    style.paragraph_format.space_after = Pt(4)

    navy = RGBColor(0x0B, 0x1F, 0x3F)
    white = RGBColor(0xFF, 0xFF, 0xFF)

    # -- En-tête --
    header = data.get("header", {})
    org = header.get("organization", "")
    if org:
        p = doc.add_paragraph(org)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in p.runs:
            run.font.size = Pt(12)
            run.font.bold = True

    title = header.get("project_title", "")
    if title:
        p = doc.add_paragraph(title)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in p.runs:
            run.font.size = Pt(10)

    qtype = header.get("questionnaire_type", "QUESTIONNAIRE")
    p = doc.add_paragraph(qtype.upper())
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in p.runs:
        run.font.size = Pt(14)
        run.font.bold = True
        run.font.color.rgb = navy

    target = header.get("target_population", "")
    zones = header.get("zones", "")
    version = header.get("version", "")
    meta_line = " — ".join(filter(None, [target, zones, version]))
    if meta_line:
        p = doc.add_paragraph(meta_line)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in p.runs:
            run.font.size = Pt(10)
            run.font.color.rgb = RGBColor(0x55, 0x55, 0x55)

    doc.add_paragraph("")

    # -- Instructions --
    instructions = data.get("instructions", [])
    if instructions:
        doc.add_heading("Instructions générales pour les enquêteurs/trices", level=1)
        for instr in instructions:
            p = doc.add_paragraph(instr, style="List Bullet")
            p.paragraph_format.space_after = Pt(2)
        doc.add_paragraph("")

    # -- Parties --
    for part in data.get("parts", []):
        part_num = part.get("number", "")
        part_title = part.get("title", "")
        doc.add_heading(f"Partie {part_num} — {part_title}", level=1)

        desc = part.get("description", "")
        if desc:
            p = doc.add_paragraph(desc)
            p.runs[0].font.italic = True if p.runs else None

        # Sous-sections
        for sub in part.get("sub_sections", []):
            sub_title = sub.get("title", "")
            if sub_title:
                doc.add_heading(sub_title, level=2)
            sub_desc = sub.get("description", "")
            if sub_desc:
                doc.add_paragraph(sub_desc)
            _add_questions_table(doc, sub.get("questions", []), navy, white)

        # Questions directes
        questions = part.get("questions", [])
        if questions:
            _add_questions_table(doc, questions, navy, white)

        doc.add_paragraph("")

    # -- Clôture --
    closure = data.get("closure_script", "")
    if closure:
        doc.add_heading("Clôture de l'entretien", level=1)
        p = doc.add_paragraph(closure)
        p.runs[0].font.italic = True if p.runs else None

    # -- Mapping indicateurs --
    mapping = data.get("indicator_mapping", [])
    if mapping:
        doc.add_heading("Correspondance questions ↔ indicateurs", level=2)
        table = doc.add_table(rows=1, cols=2)
        table.style = "Table Grid"
        hdr = table.rows[0].cells
        hdr[0].text = "Indicateur"
        hdr[1].text = "Question(s)"
        for cell in hdr:
            _shade_cell(cell, navy)
            for p in cell.paragraphs:
                for run in p.runs:
                    run.font.color.rgb = white
                    run.font.bold = True
                    run.font.size = Pt(10)
        for m in mapping:
            row = table.add_row().cells
            row[0].text = str(m.get("indicator_code", ""))
            qs = m.get("questions", [])
            row[1].text = ", ".join(str(q) for q in qs) if isinstance(qs, list) else str(qs)

    # Sauvegarder
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _add_questions_table(doc, questions, navy, white):
    """Ajouter un tableau de questions au document."""
    if not questions:
        return
    table = doc.add_table(rows=1, cols=5)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    # En-tête
    headers = ["N°", "Question", "Modalités de réponse", "Filtre / Instruction", "Indicateur mesuré"]
    hdr_cells = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr_cells[i].text = h
        _shade_cell(hdr_cells[i], navy)
        for p in hdr_cells[i].paragraphs:
            for run in p.runs:
                run.font.color.rgb = white
                run.font.bold = True
                run.font.size = Pt(9)

    # Lignes
    for q in questions:
        row_cells = table.add_row().cells
        row_cells[0].text = str(q.get("number", ""))
        row_cells[1].text = str(q.get("question", ""))
        row_cells[2].text = str(q.get("modalities", ""))
        row_cells[3].text = str(q.get("filter", "—"))
        row_cells[4].text = str(q.get("indicator", "—"))

        for cell in row_cells:
            for p in cell.paragraphs:
                for run in p.runs:
                    run.font.size = Pt(9)

    # Largeurs approximatives
    widths = [Cm(1.2), Cm(5.5), Cm(4.5), Cm(3.5), Cm(2.3)]
    for row in table.rows:
        for i, cell in enumerate(row.cells):
            cell.width = widths[i]


def _shade_cell(cell, color: RGBColor):
    """Appliquer une couleur de fond à une cellule."""
    shading = cell._element.get_or_add_tcPr()
    shd = shading.makeelement(qn("w:shd"), {
        qn("w:val"): "clear",
        qn("w:color"): "auto",
        qn("w:fill"): f"{color.red:02X}{color.green:02X}{color.blue:02X}",
    })
    # Supprimer l'ancien shading s'il existe
    for old in shading.findall(qn("w:shd")):
        shading.remove(old)
    shading.append(shd)
